import pandas as pd
import openpyxl
from datetime import datetime

pd.set_option('display.max_rows', 500)
pd.set_option('display.max_columns', 500)
pd.set_option('display.width', 1000)
pd.set_option('max_colwidth', 520)

file_1 = './excel_files/Пример №1.xlsx'
file_2 = './excel_files/Пример №2.xlsx'
file_3 = './excel_files/Пример_3_20201022.xlsx'


def get_data_from_range():
    pass


def sample_1(file, start_range='A1', end_range='B2'):
    """
    Функция получает данные из диапазона в excel файле, 'схлопывает' шапку и преобразует данные в плоскую таблицу.

    :param file: исходный excel файл
    :param start_range: стратовая ячейка диапазона
    :param end_range: конечная ячейка диапазона
    :return: итоговый DataFrame
    """
    wb = openpyxl.load_workbook(filename=file, read_only=True)
    ws = wb['Светофор №5']
    data_rows = []
    for row in ws[start_range:end_range]:
        data_cols = []
        for cell in row:
            data_cols.append(cell.value)
        data_rows.append(data_cols)
    df = pd.DataFrame(data_rows)

    # получаем шапку из столбцов для преобразования
    df_names = df.head(3)

    # получаем данные
    df_data = df.loc[4:]

    # транспонируем шапку, заполняем значения вниз, объединяем в один столбец (будущие заголовки)
    df_names = df_names.T
    df_names.ffill(axis=0, inplace=True)
    df_names.fillna('', inplace=True)
    df_names['name'] = df_names[0] + ' ' + df_names[1] + ' (' + df_names[2] + ')'
    df_names['name'].replace(to_replace=r'\n', value='', inplace=True, regex=True)

    # собираем заголовки в список, применяем новые заголовки к датафрейму с данными
    col_one_list = df_names['name'].tolist()
    df_data.columns = col_one_list

    # анпивот датафрейма с данными
    df_unpivot = pd.melt(df_data, id_vars=col_one_list[0], value_vars=col_one_list[1:],
                         var_name='Params', value_name='Values')
    df_unpivot.rename(columns={col_one_list[0]: "Subject"}, inplace=True)

    df_unpivot['time_stamp'] = datetime.now()
    return df_unpivot


def sample_3():
    wb = openpyxl.load_workbook(filename=file_2, read_only=True)
    ws = wb['Лист1']
    data_rows = []
    for row in ws['A16':'J60']:
        data_cols = []
        for cell in row:
            data_cols.append(cell.value)
        data_rows.append(data_cols)
    df = pd.DataFrame(data_rows)
    return df


if __name__ == '__main__':
    print(sample_1(file_1, start_range='B2', end_range='R78'))
    print(sample_1(file_1, start_range='B84', end_range='R96'))
